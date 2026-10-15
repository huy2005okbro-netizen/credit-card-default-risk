from pathlib import Path
from docx import Document

SOURCE = Path(r"C:\Users\trand\Downloads\30_Project_Hoc_may_co_ban_12523W.1\30_Project_Hoc_may_co_ban_12523W.1\Project_08_rui_ro_tre_han_the_tin_dung.docx")
OUTPUT = Path(r"C:\Users\trand\Documents\Codex\2026-10-12\f\outputs\Project_08_credit_card_payment_default_risk_EN.docx")

T = {
"HỌC MÁY CƠ BẢN": "FUNDAMENTALS OF MACHINE LEARNING",
"ĐỀ TÀI PROJECT KẾT THÚC HỌC PHẦN": "END-OF-COURSE PROJECT BRIEF",
"PROJECT 08": "PROJECT 08",
"SÀNG LỌC NGUY CƠ TRỄ HẠN THANH TOÁN THẺ TÍN DỤNG": "CREDIT CARD PAYMENT DEFAULT RISK SCREENING",
"Mục tiêu ngắn gọn": "Brief objective",
"Bộ phận chăm sóc khách hàng muốn phát hiện sớm tài khoản cần nhắc thanh toán hoặc tư vấn tài chính, không tự động từ chối tín dụng.": "The customer service team wants to identify early accounts that need payment reminders or financial counseling, not automatically deny credit.",
"Tài liệu giao nhiệm vụ • Phiên bản dùng cho Kết thúc học phần": "Project assignment brief • End-of-course version",
"1. Bối cảnh, mục tiêu và câu hỏi nghiên cứu": "1. Context, objectives, and research question",
"Bối cảnh thực tiễn": "Practical context",
"Câu hỏi nghiên cứu": "Research question",
"Mô hình xác suất nhận diện nhóm có nguy cơ default tháng sau tốt đến đâu khi dữ liệu mất cân bằng?": "How well can a probability model identify the group at risk of default in the following month when the data are imbalanced?",
"Mục tiêu học tập bắt buộc": "Required learning objectives",
"Vận dụng đúng kiến thức của Bài 3 — Ma trận nhầm lẫn, các chỉ số và khả năng tổng quát hóa; giải thích được giả định, ưu và nhược điểm của phương pháp.": "Correctly apply the knowledge from Lesson 3 — confusion matrices, metrics, and generalization; explain the assumptions, advantages, and disadvantages of the methods.",
"Xây dựng pipeline có thể chạy lại từ dữ liệu thô đến kết quả; giữ tập kiểm thử độc lập và báo cáo trung thực.": "Build a reproducible pipeline from raw data to results; keep an independent test set and report honestly.",
"So sánh với baseline đơn giản, phân tích lỗi thay vì chỉ nêu một con số điểm số.": "Compare against a simple baseline and analyze errors rather than reporting only a single score.",
"Đóng gói kết quả thành ứng dụng web có API, kiểm tra đầu vào và thông tin giới hạn của mô hình.": "Package the results as a web application with an API, input validation, and information about model limitations.",
"Phân công cân bằng cho hai thành viên, thể hiện qua Git history, nhật ký và phần vấn đáp cá nhân.": "Distribute work fairly between the two members, as evidenced by Git history, logs, and the individual oral examination.",
"Đặc tả bài toán học máy": "Machine learning problem specification",
"Tiêu chí thành công tối thiểu": "Minimum success criteria",
"Nhóm phải chạy được toàn bộ pipeline trên máy khác, vượt hoặc giải thích rõ khi không vượt baseline, và chứng minh việc chia tập/tiền xử lý không làm rò rỉ thông tin.": "The team must be able to run the full pipeline on another machine, outperform the baseline or clearly explain why it does not, and demonstrate that splitting and preprocessing do not leak information.",
"2. Dữ liệu thực tế và quy tắc sử dụng": "2. Real-world data and usage rules",
"Nguồn chính thức": "Official source",
"Trang dữ liệu: https://archive.ics.uci.edu/dataset/350/default%2Bof%2Bcredit%2Bcard%2Bclients": "Dataset page: https://archive.ics.uci.edu/dataset/350/default%2Bof%2Bcredit%2Bcard%2Bclients",
"DOI/tài liệu mô tả: https://doi.org/10.24432/C55S3H": "DOI/documentation: https://doi.org/10.24432/C55S3H",
"Hồ sơ dữ liệu phải nộp": "Required data documentation",
"README dữ liệu: nguồn, ngày tải, tệp sử dụng, checksum hoặc phiên bản, giấy phép/điều khoản và trích dẫn.": "Data README: source, download date, files used, checksum or version, license/terms, and citation.",
"Data dictionary: tên biến, kiểu, đơn vị, vai trò feature/target/ID, thời điểm biến trở nên có sẵn.": "Data dictionary: variable name, type, unit, feature/target/ID role, and when the variable becomes available.",
"Báo cáo chất lượng: missing, trùng lặp, ngoại lệ, phân bố lớp/target và mọi dòng bị loại kèm lý do.": "Quality report: missing values, duplicates, outliers, class/target distribution, and every removed row with its reason.",
"Không đưa dữ liệu lớn hoặc dữ liệu có hạn chế vào Git; cung cấp script tải hoặc hướng dẫn tái tạo.": "Do not put large or restricted data in Git; provide a download script or reproduction instructions.",
"Mốc thời gian và rò rỉ dữ liệu": "Timeline and data leakage",
"Chỉ dùng thông tin trước tháng mục tiêu. Preprocessing và cân bằng lớp chỉ thực hiện trong fold train; tuyệt đối không oversample trước khi chia.": "Use only information available before the target month. Perform preprocessing and class balancing only within the training fold; never oversample before splitting.",
"Quy tắc không thương lượng": "Non-negotiable rules",
"Tách train/validation/test trước mọi bước học từ dữ liệu; mọi imputer, scaler, encoder, feature selection, PCA hoặc resampling phải fit đúng phần train.": "Split train/validation/test before any data-learning step; every imputer, scaler, encoder, feature selection method, PCA, or resampling method must be fit on the appropriate training portion.",
"Test chỉ dùng một lần để báo cáo cuối. Không chọn mô hình, ngưỡng hoặc tham số dựa trên test.": "Use the test set only once for final reporting. Do not select a model, threshold, or parameter based on the test set.",
"Ghi random_state và môi trường; nếu dữ liệu theo thời gian/người/phiên/hóa đơn, chia theo đúng nhóm tương ứng.": "Record the random_state and environment; if the data are organized by time/person/session/invoice, split by the corresponding group.",
"3. Quy trình kỹ thuật và thí nghiệm bắt buộc": "3. Required technical process and experiments",
"Pipeline tối thiểu": "Minimum pipeline",
"Đóng băng định nghĩa bài toán, thời điểm dự đoán và tiêu chí đánh giá trước khi xem test.": "Freeze the problem definition, prediction time point, and evaluation criteria before viewing the test set.",
"Tải dữ liệu bằng script; kiểm tra schema, số dòng, khóa và chất lượng dữ liệu.": "Load data using a script; check the schema, row count, keys, and data quality.",
"Khám phá dữ liệu trên train; lập biểu đồ phục vụ quyết định tiền xử lý, không chỉ trang trí.": "Explore data on the training set; create plots that support preprocessing decisions, not merely decoration.",
"Xây dựng baseline: Tất cả không-default và logistic regression không class_weight.": "Build baselines: always non-default and logistic regression without class_weight.",
"Xây dựng mô hình chính: Logistic regression với/không class_weight; ngưỡng theo ma trận chi phí giả định.": "Build the main model: logistic regression with/without class_weight; set the threshold using an assumed cost matrix.",
"Chọn tham số bằng validation/CV phù hợp; lưu toàn bộ Pipeline và cấu hình.": "Select parameters using suitable validation/CV; save the entire pipeline and configuration.",
"Đánh giá một lần trên test; báo khoảng biến thiên qua seed/fold khi phù hợp.": "Evaluate once on the test set; report variation across seeds/folds when appropriate.",
"Phân tích lỗi, giới hạn, trường hợp ngoài miền và chuyển mô hình vào API/web.": "Analyze errors, limitations, and out-of-domain cases, then deploy the model to the API/web app.",
"Bốn thí nghiệm bắt buộc": "Four required experiments",
"Chuẩn so sánh": "Comparison standard",
"Dùng cùng split và cùng đơn vị đánh giá cho các mô hình.": "Use the same split and evaluation unit for all models.",
"Báo cả metric chính lẫn ít nhất một biểu đồ/ma trận lỗi phù hợp.": "Report both the primary metric and at least one appropriate plot/error matrix.",
"Tách rõ kết quả validation (dùng chọn) và test (dùng kết luận).": "Clearly separate validation results (used for selection) from test results (used for conclusions).",
"Không khẳng định quan hệ nhân quả từ hệ số, feature importance, cụm hoặc tương quan.": "Do not claim causal relationships from coefficients, feature importance, clusters, or correlations.",
"Cấu trúc mã nguồn gợi ý": "Suggested source-code structure",
"4. Sản phẩm web và hồ sơ bàn giao": "4. Web product and handover package",
"Yêu cầu ứng dụng web": "Web application requirements",
"Form mô phỏng và dashboard ngưỡng; API POST /api/risk-score; trang model card; không lưu đầu vào.": "Simulation form and threshold dashboard; POST /api/risk-score API; model card page; do not store inputs.",
"Có ít nhất 03 màn hình: giới thiệu/phạm vi; thao tác chính; dashboard đánh giá & model card.": "Include at least 03 screens: introduction/scope; main workflow; evaluation dashboard & model card.",
"API trả JSON, có schema/validation, mã lỗi hợp lý, ví dụ request/response và kiểm thử tối thiểu.": "The API returns JSON and includes a schema/validation, sensible error codes, request/response examples, and minimum tests.",
"Giao diện hiển thị đơn vị, miền giá trị và thông báo khi đầu vào thiếu/ngoài miền; không âm thầm thay thế giá trị.": "The interface shows units, valid ranges, and messages for missing/out-of-range inputs; do not silently substitute values.",
"Tách train offline khỏi serving. Ứng dụng chỉ nạp model/pipeline đã lưu, không huấn luyện lại ở mỗi request.": "Separate offline training from serving. The application loads only the saved model/pipeline and does not retrain on each request.",
"Stack khuyến nghị: Python + pandas + scikit-learn; Flask/FastAPI; HTML/CSS/JS hoặc React/Vue. SQLite chỉ khi thật sự cần.": "Recommended stack: Python + pandas + scikit-learn; Flask/FastAPI; HTML/CSS/JS or React/Vue. Use SQLite only when truly needed.",
"Các tệp/sản phẩm phải nộp": "Required files/deliverables",
"Phần mở rộng (chỉ chấm khi phần bắt buộc hoàn chỉnh)": "Extensions (graded only when the required work is complete)",
"Hiệu chỉnh xác suất hoặc phân tích độ nhạy khi bỏ biến nhạy cảm.": "Probability calibration or sensitivity analysis when sensitive variables are removed.",
"Nguyên tắc phạm vi": "Scope principles",
"Không bắt buộc cloud, mobile app, deep learning ngoài bài hoặc thu thập thêm dữ liệu. Một hệ thống nhỏ nhưng tái lập, trung thực và phân tích sâu được đánh giá cao hơn sản phẩm nhiều tính năng nhưng thiếu bằng chứng.": "Cloud deployment, a mobile app, deep learning beyond the course, and collecting additional data are not required. A small system that is reproducible, honest, and deeply analyzed is valued more highly than a feature-rich product lacking evidence.",
"5. Kế hoạch thực hiện trong 06 tuần": "5. Six-week implementation plan",
"Phân công hai thành viên": "Two-member work allocation",
"Mỗi người phải có đóng góp đáng kể ở cả phần dữ liệu/mô hình và phần web/báo cáo; không chia cứng một người chỉ viết báo cáo.": "Each person must make a substantial contribution to both the data/model work and the web/report work; do not rigidly assign one person only to write the report.",
"Pull request/commit message phải mô tả công việc. Nhật ký tuần ghi người thực hiện, giờ ước lượng, kết quả và vấn đề.": "Pull requests/commit messages must describe the work. The weekly log records the contributor, estimated hours, results, and issues.",
"Hai thành viên cùng hiểu toàn bộ pipeline và đều có thể trả lời câu hỏi về leakage, metric, mô hình và API.": "Both members understand the full pipeline and can answer questions about leakage, metrics, the model, and the API.",
"Cổng kiểm tra trước khi nộp": "Pre-submission checks",
"6. Rubric đánh giá và yêu cầu báo cáo": "6. Assessment rubric and report requirements",
"Tổng: 100 điểm. Phần mở rộng chỉ được cộng trong phạm vi tiêu chí liên quan, không bù cho sai split, leakage hoặc sản phẩm không chạy.": "Total: 100 points. Extension work may add points only within relevant criteria; it does not compensate for an incorrect split, leakage, or a non-running product.",
"Điều kiện giới hạn điểm": "Score caps",
"Không tái lập được kết quả, không có test độc lập hoặc có leakage nghiêm trọng chưa sửa: tối đa 50/100.": "If results are not reproducible, there is no independent test set, or serious leakage remains unfixed: maximum 50/100.",
"Không có ứng dụng web/API chạy được: mất toàn bộ mục 6; báo cáo không thay thế sản phẩm.": "If there is no working web application/API: lose all points for section 6; the report does not replace the product.",
"Không trích nguồn dữ liệu/vi phạm điều khoản hoặc sao chép: xử lý theo quy chế học vụ.": "Failure to cite the data source, violation of terms, or plagiarism will be handled under academic regulations.",
"Một thành viên không chứng minh được đóng góp/hiểu biết: điểm cá nhân có thể khác điểm nhóm.": "If a member cannot demonstrate contribution/understanding, the individual grade may differ from the group grade.",
"Đề cương báo cáo bắt buộc": "Required report outline",
"Tóm tắt • Bối cảnh & câu hỏi • Dữ liệu/giấy phép • Thời điểm & leakage • Phương pháp • Thiết kế thí nghiệm • Kết quả • Phân tích lỗi • Web/API • Đạo đức & giới hạn • Kết luận • Tài liệu tham khảo • Phụ lục tái lập.": "Abstract • Context & question • Data/license • Time point & leakage • Methods • Experiment design • Results • Error analysis • Web/API • Ethics & limitations • Conclusion • References • Reproducibility appendix.",
"Câu hỏi vấn đáp gợi ý": "Suggested oral-examination questions",
"Vì sao oversampling trước split gây sai?": "Why does oversampling before splitting cause errors?",
"Calibration khác discrimination thế nào?": "How does calibration differ from discrimination?",
"Một mô hình có AUC tốt vẫn có thể không công bằng ra sao?": "How can a model with good AUC still be unfair?",
"Sử dụng có trách nhiệm": "Responsible use",
"Chỉ phục vụ học tập và hỗ trợ chăm sóc; không dùng cho quyết định cấp tín dụng, xử phạt hay định giá cá nhân.": "For learning and customer-support purposes only; do not use for credit decisions, penalties, or individual pricing.",
"Tài liệu tham khảo tối thiểu": "Minimum references",
"[1] Slide học phần: Bài 3 — Ma trận nhầm lẫn, các chỉ số và khả năng tổng quát hóa.\n[2] Nguồn dữ liệu: https://archive.ics.uci.edu/dataset/350/default%2Bof%2Bcredit%2Bcard%2Bclients\n[3] DOI/mô tả: https://doi.org/10.24432/C55S3H\n[4] Tài liệu phương pháp scikit-learn: https://scikit-learn.org/stable/modules/model_evaluation.html": "[1] Course slides: Lesson 3 — Confusion matrices, metrics, and generalization.\n[2] Data source: https://archive.ics.uci.edu/dataset/350/default%2Bof%2Bcredit%2Bcard%2Bclients\n[3] DOI/documentation: https://doi.org/10.24432/C55S3H\n[4] scikit-learn methodology documentation: https://scikit-learn.org/stable/modules/model_evaluation.html",
"THÔNG TIN": "INFORMATION", "NỘI DUNG": "CONTENT", "Chủ đề học máy": "Machine learning topic", "Bài 3 — Ma trận nhầm lẫn, các chỉ số và khả năng tổng quát hóa": "Lesson 3 — Confusion matrices, metrics, and generalization", "Lớp": "Class", "Giảng viên": "Instructor", "Hình thức": "Format", "Nhóm 02 sinh viên": "Team of 02 students", "Thời gian": "Duration", "06 tuần": "06 weeks", "Mức độ": "Level", "Trung bình — có phần mở rộng tùy chọn": "Intermediate — optional extensions", "Dữ liệu": "Dataset",
"PGS.TS. NGUYỄN VĂN HẬU": "Assoc. Prof. Dr. NGUYEN VAN HAU",
"PROJECT 08  |  BÀI 3 — MA TRẬN NHẦM LẪN, CÁC CHỈ SỐ VÀ KHẢ NĂNG TỔNG QUÁT HÓA": "PROJECT 08  |  LESSON 3 — CONFUSION MATRICES, METRICS, AND GENERALIZATION",
"Thành phần": "Component", "Yêu cầu": "Requirement", "Loại bài toán": "Problem type", "Phân loại nhị phân và lựa chọn ngưỡng theo chi phí.": "Binary classification and cost-based threshold selection.", "Đơn vị quan sát": "Unit of observation", "Một khách hàng tại thời điểm chốt dữ liệu.": "One customer at the data snapshot time.", "Đầu vào": "Inputs", "Hạn mức, lịch sử trạng thái trả nợ, số tiền hóa đơn và thanh toán.": "Credit limit, repayment-status history, bill amounts, and payment amounts.", "Đầu ra": "Outputs", "Xác suất default tháng sau, nhóm rủi ro và giải thích mức hệ thống.": "Probability of default in the following month, risk group, and system-level explanations.", "Chỉ số chính": "Primary metrics", "PR-AUC, ROC-AUC, recall lớp default, precision, calibration/Brier score.": "PR-AUC, ROC-AUC, recall for the default class, precision, calibration/Brier score.",
"Thuộc tính": "Attribute", "Thông tin đã xác minh": "Verified information", "Tên dữ liệu": "Dataset name", "Quy mô": "Size", "30.000 khách hàng, 23 biến dự báo": "30,000 customers, 23 predictor variables", "Quyền sử dụng": "Usage rights", "Phạm vi đề xuất": "Recommended scope", "Giữ test độc lập; làm rõ mã PAY_*; không coi ID là đặc trưng; phân tích phân bố lớp.": "Keep the test set independent; clarify PAY_* codes; do not treat ID as a feature; analyze class distribution.",
"#": "#", "Thí nghiệm / bằng chứng cần có": "Experiment / required evidence", "So sánh accuracy với PR-AUC/recall.": "Compare accuracy with PR-AUC/recall.", "Đánh giá calibration.": "Evaluate calibration.", "So sánh 3 ngưỡng và chi phí.": "Compare 3 thresholds and costs.", "Audit sai số theo sex/education/marriage; nêu giới hạn dữ liệu.": "Audit errors by sex/education/marriage; state data limitations.",
"Hạng mục": "Item", "Yêu cầu tối thiểu": "Minimum requirement", "Kho mã nguồn": "Source repository", "Code sạch; README chạy từ máy mới; requirements/lock; commit của cả hai thành viên.": "Clean code; README runnable on a new machine; requirements/lock file; commits by both members.", "Dữ liệu": "Data", "README, data dictionary, script tải/làm sạch, không vi phạm giấy phép.": "README, data dictionary, download/cleaning scripts, no license violations.", "Mô hình": "Model", "Pipeline đã lưu, config/seed, script train và evaluate tách biệt.": "Saved pipeline, config/seed, separate train and evaluate scripts.", "Ứng dụng web": "Web application", "Chạy local; API; validation; dashboard/model card; có dữ liệu demo an toàn.": "Runs locally; API; validation; dashboard/model card; safe demo data.", "Báo cáo": "Report", "15–25 trang không tính phụ lục; PDF và tệp nguồn DOCX/LaTeX; biểu đồ đọc được.": "15–25 pages excluding appendices; PDF and DOCX/LaTeX source files; readable charts.", "Trình bày": "Presentation", "10–12 slide; demo trực tiếp 5–7 phút; tổng 12–15 phút và vấn đáp.": "10–12 slides; a 5–7 minute live demo; 12–15 minutes total plus oral examination.", "Bàn giao": "Handover", "Link repo, hướng dẫn, model/data card, nhật ký 6 tuần, bảng phân công và tự đánh giá.": "Repository link, instructions, model/data card, six-week log, work-allocation table, and self-assessment.",
"Tuần": "Week", "Công việc chính": "Main work", "Minh chứng đầu ra": "Deliverable evidence", "Chốt câu hỏi, thời điểm dự đoán, nguồn/giấy phép; tạo repo và backlog.": "Finalize the question, prediction time point, source/license; create the repository and backlog.", "Project brief 1–2 trang; data README; schema; baseline plan.": "1–2 page project brief; data README; schema; baseline plan.", "Tải/làm sạch; EDA trên train; hoàn thiện split và baseline.": "Download/clean data; EDA on the training set; finalize split and baseline.", "Script dữ liệu; notebook EDA; baseline chạy được; checkpoint GV.": "Data scripts; EDA notebook; working baseline; instructor checkpoint.", "Pipeline mô hình chính; validation/CV; theo dõi thí nghiệm.": "Main-model pipeline; validation/CV; experiment tracking.", "Bảng thí nghiệm vòng 1; model candidate; kiểm thử preprocessing.": "Round-1 experiment table; model candidate; preprocessing tests.", "Hoàn tất thí nghiệm, chọn mô hình/ngưỡng; test cuối; phân tích lỗi.": "Complete experiments, select model/threshold; final test; error analysis.", "Bảng kết quả đóng băng; figures; model card nháp.": "Frozen results table; figures; model-card draft.", "Đóng gói API và giao diện; tích hợp model; test đầu vào và lỗi.": "Package the API and interface; integrate the model; test inputs and errors.", "Web demo end-to-end; API docs; test; checkpoint demo.": "End-to-end web demo; API documentation; tests; demo checkpoint.", "Hoàn thiện báo cáo/slide; rehearsal; tái lập trên môi trường sạch.": "Complete the report/slides; rehearse; reproduce in a clean environment.", "Release cuối; báo cáo; slide; biên bản phân công.": "Final release; report; slides; work-allocation record.",
"Cổng": "Gate", "Đạt khi": "Pass when", "Nguồn truy cập được; giấy phép/trích dẫn rõ; script tái tạo; không dữ liệu nhạy cảm thật.": "Source accessible; license/citation clear; reproduction script; no real sensitive data.", "Đánh giá": "Evaluation", "Split hợp lệ; baseline; validation/test tách; bảng metric và phân tích lỗi.": "Valid split; baseline; separate validation/test; metrics table and error analysis.", "Kỹ thuật": "Technical", "Một lệnh hoặc hướng dẫn ngắn chạy train/evaluate/app; không đường dẫn tuyệt đối cá nhân.": "One command or short instructions run train/evaluate/app; no personal absolute paths.", "Web": "Web", "Luồng chính và API chạy; input xấu không làm sập; model card/giới hạn hiện rõ.": "Main workflow and API run; bad input does not crash it; model card/limitations are clearly visible.", "Học thuật": "Academic", "Mọi bảng/hình có nguồn; không sao chép; ghi rõ công cụ AI và cách kiểm chứng nếu có.": "Every table/figure has a source; no plagiarism; clearly disclose any AI tools and how their output was verified.",
"Tiêu chí": "Criterion", "Điểm": "Points", "Mô tả đạt yêu cầu": "Description of meeting requirements", "1. Bài toán & phạm vi": "1. Problem & scope", "Câu hỏi đo được; đơn vị quan sát, thời điểm, đầu ra và tiêu chí thành công rõ.": "Measurable question; clear unit of observation, time point, outputs, and success criteria.", "2. Dữ liệu & EDA": "2. Data & EDA", "Nguồn/giấy phép; data dictionary; chất lượng; EDA dẫn tới quyết định.": "Source/license; data dictionary; quality; EDA leading to decisions.", "3. Split, leakage & tái lập": "3. Split, leakage & reproducibility", "Chia đúng; Pipeline đúng; seed/version; test độc lập; không lối tắt.": "Correct splitting; correct pipeline; seed/version; independent test set; no shortcuts.", "4. Mô hình & thí nghiệm": "4. Model & experiments", "Baseline; mô hình đúng kiến thức chủ đề; thiết kế so sánh công bằng; ablation/tuning hợp lý.": "Baseline; model aligned with course knowledge; fair comparison design; sensible ablation/tuning.", "5. Đánh giá & phân tích lỗi": "5. Evaluation & error analysis", "Metric đúng; biểu đồ/ma trận; sai số theo nhóm/thời gian; kết luận không phóng đại.": "Correct metrics; plots/matrices; errors by group/time; non-exaggerated conclusions.", "6. Web/API": "6. Web/API", "Luồng end-to-end; validation; API; UX rõ; model card và cảnh báo.": "End-to-end flow; validation; API; clear UX; model card and warnings.", "7. Mã nguồn & bàn giao": "7. Source code & handover", "Cấu trúc, README, môi trường, test cơ bản, model/data artifact có thể tái tạo.": "Structure, README, environment, basic tests, reproducible model/data artifacts.", "8. Báo cáo, trình bày & nhóm": "8. Report, presentation & team", "Logic, trích dẫn, slide/demo, trả lời cá nhân, phân công và Git history.": "Logic, citations, slides/demo, individual responses, work allocation, and Git history.",
}

def translate_paragraph(p):
    original = p.text
    if original not in T:
        return
    translated = T[original]
    # Rebuild the run content while retaining paragraph-level formatting. This
    # also removes hyperlink elements, which are not consistently exposed as
    # runs by python-docx and would otherwise retain the old Vietnamese text.
    for child in list(p._p):
        if child.tag.rsplit('}', 1)[-1] != 'pPr':
            p._p.remove(child)
    p.add_run(translated)

def walk(container):
    for p in container.paragraphs:
        translate_paragraph(p)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                walk(cell)

def main():
    doc = Document(SOURCE)
    walk(doc)
    for section in doc.sections:
        walk(section.header)
        walk(section.footer)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == "__main__":
    main()
